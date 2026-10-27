"""ورود و خروجی اکسل فروش، هم‌قالب با فایل Sales CRM."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.db import transaction

from .jalali import (
    JALALI_MONTHS, gregorian_to_jalali, jalali_to_gregorian, month_key, parse_jalali, to_en_digits, to_fa_digits,
)
from .levels import level_for
from .models import Check, Customer, Invoice, InvoiceLine, Operator, Payment
from .perms import set_membership

GENERAL_EXPERT = 'عمومی'
SHEETS = ('Sales', 'Customers', 'Sale_Items', 'Payments', 'Checks', 'Commissions')
MAX_BYTES = 8 * 1024 * 1024

HEADERS = {
    'Sales': (
        'شماره فاکتور ', 'تاریخ فروش ', 'کارشناس فروش', 'نام مشتری ', 'شماره تماس ', 'شهر ', 'آدرس ',
        'نوع تسویه', 'مبلغ کل فروش ', 'مبلغ کل خرید ', 'سود کل ', 'جمع دریافتی ', 'مانده حساب ',
        'وضعیت مالی', 'توضیحات',
    ),
    'Customers': (
        'نام مشتری', 'شماره تماس', 'شهر', 'آدرس', 'تاریخ اولین خرید', 'نوع مشتری',
    ),
    'Sale_Items': (
        'شماره فاکتور ', 'نام محصول ', 'تعداد ', 'قیمت فروش ', 'قیمت خرید ',
        'مبلغ کل فروش ', 'مبلغ کل خرید ', 'سود ',
    ),
    'Payments': (
        'شماره فاکتور ', 'نوع پرداخت', 'مبلغ', 'تاریخ پرداخت ', 'توضیحات ',
    ),
    'Checks': (
        'شماره فاکتور', 'مبلغ', 'تاریخ سر رسید', 'وضعیت چک', 'توضیحات',
    ),
    'Commissions': (
        'شماره فاکتور ', 'کارشناس فروش', 'سود فاکتور', 'درصد پورسانت', 'مبلغ پورسانت', 'وضعیت پرداخت ',
    ),
}

MONTH_INDEX = {name: index + 1 for index, name in enumerate(JALALI_MONTHS)}


class WorkbookError(Exception):
    """خطای قابل‌نمایش برای فایل اکسل."""


def _text(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.isoformat()
    return ' '.join(str(value or '').replace('\xa0', ' ').split())


def _key(value):
    return _text(value).replace('\u200c', '').casefold()


def _phone(value):
    if value is None or value == '':
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = ''.join(ch for ch in to_en_digits(value) if ch.isdigit())
    if len(digits) == 10 and digits.startswith('9'):
        digits = '0' + digits
    elif len(digits) == 12 and digits.startswith('98'):
        digits = '0' + digits[2:]
    return digits[:20]


def _money(value):
    if value is None or value == '':
        return 0
    if isinstance(value, str) and value.strip() in {'?', '-', '—'}:
        return 0
    return int(round(float(value)))


def _invoice_number(value):
    if isinstance(value, bool) or value is None or _text(value) == '':
        raise WorkbookError('شماره فاکتور خالی است.')
    if isinstance(value, float) and value.is_integer():
        number = int(value)
    elif isinstance(value, int):
        number = value
    else:
        raw = to_en_digits(value).strip()
        if not raw.isdigit():
            raise WorkbookError(f'شماره فاکتور باید عدد باشد: {value!r}')
        number = int(raw)
    if number <= 0:
        raise WorkbookError(f'شماره فاکتور نامعتبر است: {number}')
    return str(number)


def _paid(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return str(value).strip().casefold() in {'1', 'true', 'yes', 'بله', 'پرداخت شده', 'پرداخت‌شده'}


def _qty(value):
    number = Decimal(str(value or 0))
    if number == number.to_integral():
        return int(number)
    return float(number)


def _month_label(value):
    if not value:
        return ''
    jy, jm, _jd = gregorian_to_jalali(value.year, value.month, value.day)
    return f'{JALALI_MONTHS[jm - 1]} {to_fa_digits(f"{jy % 100:02d}")}'


def _jalali_text(value):
    if not value:
        return ''
    jy, jm, jd = gregorian_to_jalali(value.year, value.month, value.day)
    return f'{jy:04d}/{jm:02d}/{jd:02d}'


def _sold_on(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_jalali(value)
    if parsed:
        return parsed
    parts = to_en_digits(value).split()
    if len(parts) < 2 or parts[0] not in MONTH_INDEX:
        raise WorkbookError(f'تاریخ فروش نامعتبر است: {value!r}')
    year = int(parts[-1])
    if year < 100:
        year += 1400
    month = MONTH_INDEX[parts[0]]
    gy, gm, gd = jalali_to_gregorian(year, month, 1)
    return date(gy, gm, gd)


def _exact_date(value, label):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_jalali(value)
    if parsed is None:
        raise WorkbookError(f'{label} نامعتبر است: {value!r}')
    return parsed


def _rows(sheet, width):
    found = []
    for row in sheet.iter_rows(min_row=2, max_col=width, values_only=True):
        cells = (list(row) + [None] * width)[:width]
        if all(cell is None or str(cell).strip() == '' for cell in cells):
            continue
        found.append(cells)
    return found


def _load(source):
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise WorkbookError('بسته openpyxl نصب نیست.') from exc
    try:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise WorkbookError('فایل اکسل خوانده نشد. فقط xlsx با همین قالب پذیرفته می‌شود.') from exc
    try:
        missing = set(SHEETS) - set(workbook.sheetnames)
        if missing:
            raise WorkbookError('این شیت‌ها در فایل نیستند: ' + '، '.join(sorted(missing)))
        widths = {'Sales': 15, 'Customers': 6, 'Sale_Items': 8, 'Payments': 5, 'Checks': 5, 'Commissions': 6}
        return {
            'sales': _rows(workbook['Sales'], widths['Sales']),
            'customers': _rows(workbook['Customers'], widths['Customers']),
            'items': _rows(workbook['Sale_Items'], widths['Sale_Items']),
            'payments': _rows(workbook['Payments'], widths['Payments']),
            'checks': _rows(workbook['Checks'], widths['Checks']),
            'commissions': _rows(workbook['Commissions'], widths['Commissions']),
        }
    finally:
        workbook.close()


def _parse(book):
    sales = []
    seen = set()
    for row in book['sales']:
        number = _invoice_number(row[0])
        if number in seen:
            raise WorkbookError(f'شماره فاکتور {number} در شیت فروش تکراری است.')
        seen.add(number)
        customer = _text(row[3])
        if not customer:
            raise WorkbookError(f'نام مشتری فاکتور {number} خالی است.')
        sales.append({
            'number': number,
            'sold_on': _sold_on(row[1]),
            'expert': _text(row[2]),
            'customer': customer,
            'phone': _phone(row[4]),
            'city': _text(row[5])[:80],
            'address': _text(row[6]),
            'settlement': _text(row[7]) or 'نقدی',
            'purchase': _money(row[9]),
            'status': _text(row[13]) or 'تکمیل شده',
            'note': _text(row[14]),
        })
    known = {row['number'] for row in sales}
    customers = {}
    for row in book['customers']:
        name = _text(row[0])
        key = _key(name)
        if not key:
            continue
        customers[key] = {
            'name': name[:160],
            'phone': _phone(row[1]),
            'city': _text(row[2])[:80],
            'address': _text(row[3]),
            'customer_type': _text(row[5])[:40],
        }
    items = defaultdict(list)
    for row in book['items']:
        number = _invoice_number(row[0])
        if number not in known:
            raise WorkbookError(f'فاکتور {number} در شیت اقلام هست ولی در شیت فروش نیست.')
        items[number].append({
            'product_name': _text(row[1])[:160] or 'کالا',
            'qty': Decimal(str(row[2] or 0)),
            'unit_price': _money(row[3]),
            'unit_cost': _money(row[4]),
        })
    payments = defaultdict(list)
    for row in book['payments']:
        number = _invoice_number(row[0])
        if number not in known:
            raise WorkbookError(f'فاکتور {number} در شیت پرداخت هست ولی در شیت فروش نیست.')
        raw_date = row[3]
        paid_on = None
        note = _text(row[4])
        if _text(raw_date):
            try:
                paid_on = _exact_date(raw_date, 'تاریخ پرداخت')
            except WorkbookError:
                note = (note + ' ' if note else '') + f'تاریخ پرداخت در فایل: {_text(raw_date)}'
        payments[number].append({
            'method': _text(row[1])[:80] or 'شبا',
            'amount': _money(row[2]),
            'paid_on': paid_on,
            'note': note[:240],
        })
    checks = defaultdict(list)
    for row in book['checks']:
        number = _invoice_number(row[0])
        if number not in known:
            raise WorkbookError(f'فاکتور {number} در شیت چک هست ولی در شیت فروش نیست.')
        checks[number].append({
            'amount': _money(row[1]),
            'due_on': _exact_date(row[2], f'سررسید چک فاکتور {number}'),
            'status': _text(row[3])[:40] or 'در انتظار',
            'note': _text(row[4])[:240],
        })
    commissions = {}
    for row in book['commissions']:
        number = _invoice_number(row[0])
        if number not in known:
            raise WorkbookError(f'فاکتور {number} در شیت پورسانت هست ولی در شیت فروش نیست.')
        commissions[number] = _paid(row[5])
    if not sales and not customers:
        raise WorkbookError('فایل ردیف قابل ورود ندارد.')
    return {
        'sales': sales,
        'customers': customers,
        'items': items,
        'payments': payments,
        'checks': checks,
        'commissions': commissions,
    }


def _experts(parsed):
    found = {}
    for user in Operator.objects.all():
        for label in (user.expert_name, user.full_name, user.username):
            key = _key(label)
            if key and key not in found:
                found[key] = user
    created = 0
    for row in parsed['sales']:
        name = row['expert']
        key = _key(name)
        if not key or key == _key(GENERAL_EXPERT) or key in found:
            continue
        user = Operator(username=name[:150], full_name=name[:120], expert_name=name[:120], is_active=True)
        user.set_unusable_password()
        user.save()
        set_membership(user, 'specialist', False)
        found[key] = user
        created += 1
    return found, created


def _fill_customer(record, row):
    if not record['phone']:
        record['phone'] = row['phone']
    if not record['city']:
        record['city'] = row['city']
    if not record['address']:
        record['address'] = row['address']


def _customers(parsed):
    records = {key: dict(row) for key, row in parsed['customers'].items()}
    for row in parsed['sales']:
        key = _key(row['customer'])
        record = records.get(key)
        if record is None:
            record = {
                'name': row['customer'][:160],
                'phone': '',
                'city': '',
                'address': '',
                'customer_type': '',
            }
            records[key] = record
        _fill_customer(record, row)
    existing = {}
    for customer in Customer.objects.all():
        key = _key(customer.name)
        if key and key not in existing:
            existing[key] = customer
    created = 0
    touched = []
    for key, record in records.items():
        customer = existing.get(key)
        if customer is None:
            customer = Customer.objects.create(**record)
            existing[key] = customer
            created += 1
        else:
            changed = []
            for field in ('phone', 'city', 'address', 'customer_type'):
                value = record[field]
                if value and not getattr(customer, field):
                    setattr(customer, field, value)
                    changed.append(field)
            if changed:
                customer.save(update_fields=changed)
        touched.append(customer)
    return existing, created, touched


def _assign_experts(customers):
    for customer in customers:
        if customer.expert_id:
            continue
        counts = defaultdict(int)
        for invoice in customer.invoices.all():
            if invoice.expert_id:
                counts[invoice.expert_id] += 1
        if not counts:
            continue
        customer.expert_id = max(counts, key=lambda expert_id: (counts[expert_id], expert_id))
        customer.save(update_fields=['expert'])


def import_workbook(source):
    """فاکتور و مشتری فایل را اضافه می‌کند. فاکتور هم‌شماره به‌روز می‌شود."""
    if hasattr(source, 'size') and source.size > MAX_BYTES:
        raise WorkbookError('حجم فایل بیشتر از ۸ مگابایت است.')
    raw = source.read() if hasattr(source, 'read') else source
    if isinstance(raw, (bytes, bytearray)) and len(raw) > MAX_BYTES:
        raise WorkbookError('حجم فایل بیشتر از ۸ مگابایت است.')
    parsed = _parse(_load(BytesIO(raw) if isinstance(raw, (bytes, bytearray)) else raw))
    with transaction.atomic():
        experts, experts_created = _experts(parsed)
        customers, customers_created, touched = _customers(parsed)
        fresh = updated = 0
        for row in parsed['sales']:
            expert_key = _key(row['expert'])
            expert = None if expert_key in {'', _key(GENERAL_EXPERT)} else experts[expert_key]
            customer = customers[_key(row['customer'])]
            invoice = Invoice.objects.filter(number=row['number']).first()
            fields = {
                'customer': customer,
                'expert': expert,
                'sold_on': row['sold_on'],
                'settlement': row['settlement'][:20],
                'status': row['status'][:40],
                'note': row['note'],
                'purchase_amount': row['purchase'],
            }
            if invoice is None:
                invoice = Invoice.objects.create(number=row['number'], **fields)
                fresh += 1
            else:
                for name, value in fields.items():
                    setattr(invoice, name, value)
                invoice.save(update_fields=list(fields))
                updated += 1
            number = row['number']
            if number in parsed['items']:
                invoice.lines.all().delete()
                InvoiceLine.objects.bulk_create([
                    InvoiceLine(invoice=invoice, **line) for line in parsed['items'][number]
                ])
            if number in parsed['payments']:
                invoice.payments.all().delete()
                Payment.objects.bulk_create([
                    Payment(
                        invoice=invoice,
                        amount=pay['amount'],
                        method=pay['method'],
                        paid_on=pay['paid_on'] or invoice.sold_on,
                        note=pay['note'],
                    )
                    for pay in parsed['payments'][number]
                ])
            if number in parsed['checks']:
                invoice.checks.all().delete()
                Check.objects.bulk_create([
                    Check(invoice=invoice, **check) for check in parsed['checks'][number]
                ])
            if number in parsed['commissions']:
                invoice.commission_paid = parsed['commissions'][number]
                invoice.save(update_fields=['commission_paid'])
        _assign_experts(touched)
    return {
        'invoices_created': fresh,
        'invoices_updated': updated,
        'customers_created': customers_created,
        'customers': len(touched),
        'experts_created': experts_created,
    }


def _style(workbook):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color='FFFFFF', name='Calibri')
    header_fill = PatternFill('solid', fgColor='0F172A')
    for sheet in workbook.worksheets:
        sheet.sheet_view.rightToLeft = True
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='right')
        for index, column in enumerate(sheet.columns, start=1):
            width = 14
            for cell in column:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, 42))
            sheet.column_dimensions[get_column_letter(index)].width = width


def _build(rows_by_sheet):
    import openpyxl

    workbook = openpyxl.Workbook()
    first = True
    for name in SHEETS:
        sheet = workbook.active if first else workbook.create_sheet(name)
        if first:
            sheet.title = name
            first = False
        sheet.append(list(HEADERS[name]))
        for row in rows_by_sheet.get(name, ()):
            sheet.append(list(row))
    _style(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template_workbook():
    return _build({})


def _number_key(invoice):
    return int(invoice.number) if str(invoice.number).isdigit() else 10**12


def export_workbook():
    invoices = list(
        Invoice.objects.select_related('customer', 'expert').prefetch_related('lines', 'payments', 'checks')
    )
    invoices.sort(key=lambda invoice: (_number_key(invoice), invoice.pk))
    monthly = defaultdict(int)
    for invoice in invoices:
        if invoice.expert_id and invoice.status != 'لغو شده' and invoice.sold_on:
            monthly[(invoice.expert_id, month_key(invoice.sold_on))] += invoice.profit
    first_buy = {}
    for invoice in invoices:
        current = first_buy.get(invoice.customer_id)
        if invoice.sold_on and (current is None or invoice.sold_on < current):
            first_buy[invoice.customer_id] = invoice.sold_on
    sales_rows = []
    item_rows = []
    payment_rows = []
    check_rows = []
    commission_rows = []
    for invoice in invoices:
        number = int(invoice.number) if str(invoice.number).isdigit() else invoice.number
        sale = invoice.sale_amount
        purchase = int(invoice.purchase_amount or 0) or invoice.cost_amount
        received = sum(pay.amount for pay in invoice.payments.all())
        sales_rows.append((
            number,
            _month_label(invoice.sold_on),
            invoice.expert_label,
            invoice.customer.name,
            invoice.customer.phone,
            invoice.customer.city,
            invoice.customer.address,
            invoice.settlement,
            sale,
            purchase,
            sale - purchase,
            received,
            max(0, sale - received),
            invoice.status,
            invoice.note,
        ))
        for line in invoice.lines.all():
            qty = _qty(line.qty)
            line_sale = int(line.qty * line.unit_price)
            line_cost = int(line.qty * line.unit_cost)
            item_rows.append((
                number, line.product_name, qty, line.unit_price, line.unit_cost,
                line_sale, line_cost, line_sale - line_cost,
            ))
        for pay in invoice.payments.all():
            payment_rows.append((number, pay.method, pay.amount, _jalali_text(pay.paid_on), pay.note))
        for check in invoice.checks.all():
            check_rows.append((number, check.amount, _jalali_text(check.due_on), check.status, check.note))
        if invoice.status != 'لغو شده' and invoice.sold_on:
            level = level_for(monthly[(invoice.expert_id, month_key(invoice.sold_on))]) if invoice.expert_id else None
            rate = level['rate'] if level else 0
            profit = invoice.profit
            commission_rows.append((
                number, invoice.expert_label, profit, rate, int(profit * rate), bool(invoice.commission_paid),
            ))
    customer_rows = []
    for customer in Customer.objects.all():
        customer_rows.append((
            customer.name,
            customer.phone,
            customer.city,
            customer.address,
            _month_label(first_buy.get(customer.pk)),
            customer.customer_type,
        ))
    return _build({
        'Sales': sales_rows,
        'Customers': customer_rows,
        'Sale_Items': item_rows,
        'Payments': payment_rows,
        'Checks': check_rows,
        'Commissions': commission_rows,
    })
