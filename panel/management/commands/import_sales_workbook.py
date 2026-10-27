"""جایگزینی داده‌های فروش با فایل اکسل کشت‌یار."""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from panel.jalali import JALALI_MONTHS, jalali_to_gregorian, parse_jalali, to_en_digits
from panel.models import (
    Check, Customer, Invoice, InvoiceLine, Operator, Payment,
    SaleRequest, Snapshot,
)
from panel.perms import set_membership

GENERAL_EXPERT = 'عمومی'
MONTH_INDEX = {name: index + 1 for index, name in enumerate(JALALI_MONTHS)}


def _text(value):
    return ' '.join(str(value or '').replace('\xa0', ' ').split())


def _key(value):
    return _text(value).replace('\u200c', '').casefold()


def _phone(value):
    if value is None or value == '':
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = ''.join(ch for ch in to_en_digits(value) if ch.isdigit())
    if len(digits) == 10 and digits.startswith('9'):
        digits = '0' + digits
    elif len(digits) == 12 and digits.startswith('98'):
        digits = '0' + digits[2:]
    return digits[:20]


def _money(value):
    if value is None or value == '':
        return 0
    if isinstance(value, str) and value.strip() in {'?', '-', '—'}:
        return 0
    return int(round(float(value)))


def _month_start(value):
    parts = to_en_digits(value).split()
    if len(parts) < 2 or parts[0] not in MONTH_INDEX:
        raise CommandError(f'تاریخ ماه نامعتبر است: {value!r}')
    year = int(parts[-1])
    if year < 100:
        year += 1400
    month = MONTH_INDEX[parts[0]]
    gy, gm, gd = jalali_to_gregorian(year, month, 1)
    from datetime import date
    return date(gy, gm, gd)


def _paid(value):
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {'1', 'true', 'yes', 'بله'}


def _rows(sheet, width):
    found = []
    for row in sheet.iter_rows(min_row=2, max_col=width, values_only=True):
        if all(cell is None or str(cell).strip() == '' for cell in row):
            continue
        found.append(row)
    return found


class Command(BaseCommand):
    help = 'داده‌های فروش، مشتری، پرداخت، چک و پورسانت را از اکسل جایگزین می‌کند.'

    def add_arguments(self, parser):
        parser.add_argument('path')
        parser.add_argument('--reset', action='store_true', help='داده‌های فعلی فروش پاک شود.')

    def handle(self, *args, **options):
        path = Path(options['path'])
        if not path.is_file():
            raise CommandError(f'فایل پیدا نشد: {path}')
        if not options['reset']:
            raise CommandError('برای پاک کردن داده‌های فعلی و ورود اکسل، --reset را بدهید.')
        book = self._read(path)
        with transaction.atomic():
            kept = self._wipe()
            experts = self._experts(book)
            customers = self._customers(book, experts)
            invoices = self._invoices(book, customers, experts)
            self._lines(book, invoices)
            self._payments(book, invoices)
            self._checks(book, invoices)
            self._commissions(book, invoices)
            self._assign_customer_experts(customers.values())
        self.stdout.write(self.style.SUCCESS(
            f'وارد شد: {len(invoices)} فاکتور، {len(customers)} مشتری، {len(experts)} کارشناس.'
        ))
        self.stdout.write('حساب‌های مدیر حفظ شد: ' + '، '.join(kept))

    def _read(self, path):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError('بسته openpyxl نصب نیست.') from exc
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        missing = {'Sales', 'Customers', 'Sale_Items', 'Payments', 'Checks', 'Commissions'} - set(workbook.sheetnames)
        if missing:
            raise CommandError('این شیت‌ها در فایل نیستند: ' + '، '.join(sorted(missing)))
        book = {
            'sales': _rows(workbook['Sales'], 15),
            'customers': _rows(workbook['Customers'], 6),
            'items': _rows(workbook['Sale_Items'], 8),
            'payments': _rows(workbook['Payments'], 5),
            'checks': _rows(workbook['Checks'], 5),
            'commissions': _rows(workbook['Commissions'], 6),
        }
        workbook.close()
        numbers = [_money(row[0]) for row in book['sales']]
        if len(numbers) != len(set(numbers)) or any(number <= 0 for number in numbers):
            raise CommandError('شماره فاکتورها تکراری یا خالی است.')
        known = set(numbers)
        for label, rows in (
            ('Sale_Items', book['items']),
            ('Payments', book['payments']),
            ('Checks', book['checks']),
            ('Commissions', book['commissions']),
        ):
            unknown = {_money(row[0]) for row in rows} - known
            if unknown:
                raise CommandError(f'فاکتورهای ناشناس در {label}: {sorted(unknown)[:8]}')
        return book

    def _wipe(self):
        keep = list(Operator.objects.filter(groups__name__in=['superadmin', 'manager']).distinct())
        if not keep:
            raise CommandError('هیچ حساب مدیر برای حفظ ورود پیدا نشد.')
        SaleRequest.objects.all().delete()
        Invoice.objects.all().delete()
        Customer.objects.all().delete()
        Snapshot.objects.all().delete()
        Operator.objects.exclude(pk__in=[user.pk for user in keep]).delete()
        return [user.username for user in keep]

    def _experts(self, book):
        names = []
        seen = set()
        for row in book['sales']:
            name = _text(row[2])
            key = _key(name)
            if not key or key == _key(GENERAL_EXPERT) or key in seen:
                continue
            seen.add(key)
            names.append(name)
        experts = {}
        for name in names:
            user = Operator(username=name, full_name=name, expert_name=name, is_active=True)
            user.set_unusable_password()
            user.save()
            set_membership(user, 'specialist', False)
            experts[_key(name)] = user
        return experts

    def _customers(self, book, experts):
        records = {}
        for row in book['customers']:
            name = _text(row[0])
            key = _key(name)
            if not key:
                continue
            records[key] = {
                'name': name[:160],
                'phone': _phone(row[1]),
                'city': _text(row[2])[:80],
                'address': _text(row[3]),
                'customer_type': _text(row[5])[:40],
            }
        for row in book['sales']:
            key = _key(row[3])
            record = records.get(key)
            if record is None:
                record = {
                    'name': _text(row[3])[:160],
                    'phone': '',
                    'city': '',
                    'address': '',
                    'customer_type': '',
                }
                records[key] = record
            if not record['phone']:
                record['phone'] = _phone(row[4])
            if not record['city']:
                record['city'] = _text(row[5])[:80]
            if not record['address']:
                record['address'] = _text(row[6])
        created = {}
        for key, record in records.items():
            created[key] = Customer.objects.create(**record)
        return created

    def _invoices(self, book, customers, experts):
        invoices = {}
        for row in book['sales']:
            expert_name = _text(row[2])
            expert = None if _key(expert_name) in {'', _key(GENERAL_EXPERT)} else experts[_key(expert_name)]
            number = str(_money(row[0]))
            invoice = Invoice.objects.create(
                number=number,
                customer=customers[_key(row[3])],
                expert=expert,
                sold_on=_month_start(row[1]),
                settlement=_text(row[7]) or 'نقدی',
                status=_text(row[13]) or 'تکمیل شده',
                note=_text(row[14]),
                purchase_amount=_money(row[9]),
            )
            invoices[number] = invoice
        return invoices

    def _lines(self, book, invoices):
        for row in book['items']:
            InvoiceLine.objects.create(
                invoice=invoices[str(_money(row[0]))],
                product_name=_text(row[1])[:160] or 'کالا',
                qty=Decimal(str(row[2] or 0)),
                unit_price=_money(row[3]),
                unit_cost=_money(row[4]),
            )

    def _payments(self, book, invoices):
        for row in book['payments']:
            invoice = invoices[str(_money(row[0]))]
            raw_date = _text(row[3])
            paid_on = parse_jalali(raw_date)
            note = _text(row[4])
            if paid_on is None:
                paid_on = invoice.sold_on
                if raw_date:
                    note = (note + ' ' if note else '') + f'تاریخ پرداخت در فایل: {raw_date}'
            Payment.objects.create(
                invoice=invoice,
                amount=_money(row[2]),
                method=_text(row[1])[:80] or 'شبا',
                paid_on=paid_on,
                note=note[:240],
            )

    def _checks(self, book, invoices):
        for row in book['checks']:
            due_on = parse_jalali(_text(row[2]))
            if due_on is None:
                raise CommandError(f'سررسید چک فاکتور {row[0]} نامعتبر است: {row[2]!r}')
            Check.objects.create(
                invoice=invoices[str(_money(row[0]))],
                amount=_money(row[1]),
                due_on=due_on,
                status=_text(row[3])[:40] or 'در انتظار',
                note=_text(row[4])[:240],
            )

    def _commissions(self, book, invoices):
        for row in book['commissions']:
            invoice = invoices[str(_money(row[0]))]
            invoice.commission_paid = _paid(row[5])
            invoice.save(update_fields=['commission_paid'])

    def _assign_customer_experts(self, customers):
        for customer in customers:
            invoices = list(customer.invoices.select_related('expert'))
            counts = defaultdict(int)
            for invoice in invoices:
                if invoice.expert_id:
                    counts[invoice.expert_id] += 1
            if not counts:
                continue
            best = max(counts, key=lambda expert_id: (counts[expert_id], expert_id))
            customer.expert_id = best
            customer.save(update_fields=['expert'])
